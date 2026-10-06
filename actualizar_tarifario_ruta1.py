import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import unicodedata
import re
import urllib.request
import csv
import io
import sys
sys.stdout.reconfigure(encoding='utf-8')

def clean_text(s):
    if s is None:
        return ''
    s = str(s).strip()
    s = s.replace('hist\ufffdrico', 'histórico').replace('Hist\ufffdrico', 'Histórico')
    s = s.replace('ca\ufffda', 'caña').replace('Ca\ufffda', 'Caña').replace('CA\ufffdA', 'CAÑA')
    s = s.replace('le\ufffda', 'leña').replace('Le\ufffda', 'Leña').replace('LE\ufffdA', 'LEÑA')
    s = s.replace('s\ufffdbado', 'sábado').replace('S\ufffdbado', 'Sábado')
    s = s.replace('mill\ufffdn', 'millán').replace('Mill\ufffdn', 'Millán')
    s = s.replace('jamund\ufffd', 'jamundí').replace('Jamund\ufffd', 'Jamundí')
    s = s.replace('guacar\ufffd', 'guacarí').replace('Guacar\ufffd', 'Guacarí')
    s = s.replace('r\ufffdo', 'río').replace('R\ufffdo', 'Río')
    s = s.replace('berl\ufffdn', 'berlín').replace('Berl\ufffdn', 'Berlín')
    s = s.replace('tulu\ufffd', 'tuluá').replace('Tulu\ufffd', 'Tuluá')
    s = s.replace('jes\ufffds', 'jesús').replace('Jes\ufffds', 'Jesús')
    s = s.replace('jos\ufffd', 'josé').replace('Jos\ufffd', 'José')
    s = s.replace('rep\ufffdblica', 'república').replace('Rep\ufffdblica', 'República')
    s = s.replace('rebaj\ufffdn', 'rebajón').replace('Rebaj\ufffdn', 'Rebajón')
    s = s.replace('fabi\ufffdn', 'fabián').replace('Fabi\ufffdn', 'Fabián')
    s = s.replace('l\ufffdpez', 'lópez').replace('L\ufffdpez', 'López')
    s = s.replace('adri\ufffdn', 'adrián').replace('Adri\ufffdn', 'Adrián')
    s = s.replace('uni\ufffdn', 'unión').replace('Uni\ufffdn', 'Unión')
    s = s.replace('mu\ufffdoz', 'muñoz').replace('Mu\ufffdoz', 'Muñoz')
    s = s.replace('y\ufffdnifer', 'yénifer').replace('Y\ufffdnifer', 'Yénifer')
    s = s.replace('d\ufffdaz', 'díaz').replace('D\ufffdaz', 'Díaz')
    s = s.replace('frigor\uffdfico', 'frigorífico').replace('Frigor\uffdfico', 'Frigorífico')
    s = s.replace('galer\ufffda', 'galería').replace('Galer\ufffda', 'Galería')
    s = s.replace('bol\ufffdvar', 'bolívar').replace('Bol\ufffdvar', 'Bolívar')
    s = s.replace('guacan\ufffda', 'guacandá').replace('Guacan\ufffda', 'Guacandá')
    s = s.replace('monta\ufffda', 'montaña').replace('Monta\ufffda', 'Montaña')
    s = s.replace('\ufffd', '')
    return s

def norm(s):
    if not s: return ''
    s = clean_text(s)
    nfkd = unicodedata.normalize('NFKD', s)
    res = ''.join([c for c in nfkd if not unicodedata.combining(c)])
    return ' '.join(re.sub(r'[^a-zA-Z0-9]', ' ', res).upper().split())

print("=== 1. DESCARGANDO DATOS EN VIVO DE GOOGLE SHEETS ===")
sheet_id = '1eQSRvG7vWkIoW3AT5e6Ahi7ndWF6P4OG_Alxo2Go0lU'

def fetch_gviz(sheet_name):
    url = f'https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&sheet={sheet_name}&headers=1'
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req) as resp:
        content = resp.read().decode('utf-8')
    return list(csv.reader(io.StringIO(content)))

puntos_raw = fetch_gviz('Puntos_Rutas')
tarifas_raw = fetch_gviz('Tarifas')
prods_raw = fetch_gviz('Productos')

print(f"Puntos_Rutas cargados: {len(puntos_raw)-1} filas")
print(f"Tarifas cargadas: {len(tarifas_raw)-1} filas")
print(f"Productos cargados: {len(prods_raw)-1} filas")

# 1. Extraer Productos Activos
prods_activos = []
for r in prods_raw[1:]:
    if len(r) >= 2:
        p_name = clean_text(r[0]).upper()
        p_est = clean_text(r[1]).lower()
        if p_name and p_est != 'inactivo':
            prods_activos.append(p_name)
prods_activos = sorted(list(set(prods_activos)))
print(f"Total productos activos: {len(prods_activos)}")

# 2. Extraer Puntos_Rutas (126 puntos oficiales)
puntos_oficiales = []
for r in puntos_raw[1:]:
    if len(r) >= 3 and r[0].strip() and r[1].strip():
        puntos_oficiales.append({
            'ruta': clean_text(r[0]),
            'prov': clean_text(r[1]),
            'pto': clean_text(r[2]),
            'dir': clean_text(r[3]) if len(r) > 3 else '',
            'tel': clean_text(r[4]) if len(r) > 4 else '',
            'hora': clean_text(r[5]) if len(r) > 5 else '',
            'frec': clean_text(r[6]) if len(r) > 6 else '',
            'estado': clean_text(r[7]) if len(r) > 7 else 'Activo',
            'materias': clean_text(r[8]) if len(r) > 8 else ''
        })

print(f"Puntos_Rutas oficiales procesados: {len(puntos_oficiales)}")

# 3. Mapear Tarifas Existentes de Google Sheets
# Extraer precios de Santa Elena y Cavasa
tarifas_santa_elena = {} # prod -> {precio, estado, obs}
tarifas_cavasa = {}      # prod -> {precio, estado, obs}
otras_tarifas_existentes = []

for r in tarifas_raw[1:]:
    if len(r) < 5: continue
    ruta = clean_text(r[0])
    prov = clean_text(r[1])
    pto = clean_text(r[2])
    prod = clean_text(r[3]).upper()
    precio_str = r[4].strip()
    estado = clean_text(r[5]) if len(r) > 5 else 'Activo'
    obs = clean_text(r[6]) if len(r) > 6 else ''
    
    # Limpiar precio numérico
    precio_clean = precio_str.replace('$', '').replace('.', '').replace(',', '').strip()
    precio_num = None
    if precio_clean and precio_clean.isdigit():
        precio_num = int(precio_clean)
    elif precio_str and precio_str not in ['', 'Pendiente Precio', '0']:
        try:
            precio_num = float(precio_str)
            if precio_num.is_integer(): precio_num = int(precio_num)
        except:
            precio_num = None

    if 'Ruta 1' in ruta or 'Santa Elena' in ruta:
        if norm(prov) == 'SANTA ELENA':
            tarifas_santa_elena[norm(prod)] = {
                'prod_original': prod,
                'precio': precio_num,
                'estado': estado,
                'obs': obs
            }
        elif norm(prov) == 'CAVASA':
            tarifas_cavasa[norm(prod)] = {
                'prod_original': prod,
                'precio': precio_num,
                'estado': estado,
                'obs': obs
            }
        else:
            # Por si ya existía alguna fila de otro proveedor
            otras_tarifas_existentes.append({
                'ruta': ruta, 'prov': prov, 'pto': pto, 'prod': prod,
                'precio': precio_num, 'estado': estado, 'obs': obs
            })
    else:
        otras_tarifas_existentes.append({
            'ruta': ruta, 'prov': prov, 'pto': pto, 'prod': prod,
            'precio': precio_num, 'estado': estado, 'obs': obs
        })

print(f"Tarifas base extraídas de Santa Elena: {len(tarifas_santa_elena)} productos")
print(f"Tarifas base extraídas de Cavasa: {len(tarifas_cavasa)} productos")
print(f"Tarifas de otras rutas preservadas: {len(otras_tarifas_existentes)} filas")

# 4. Generar Nuevas Tarifas Depuradas para Google Sheets (Opción 1)
nuevas_tarifas_sheets = []
ruta1_nombre = "Ruta 1: Santa Elena / Cavasa"

# 4.1 Proveedores de Santa Elena (3)
provs_santa_elena = [
    {'prov': 'Bodega Santa Elena', 'pto': 'Santa Elena'},
    {'prov': 'Alejandro Garay', 'pto': 'Santa Elena'},
    {'prov': 'Sevillana Santa Elena', 'pto': 'Sevillana Santa Elena'}
]

for p in provs_santa_elena:
    for prod_n, item in sorted(tarifas_santa_elena.items()):
        nuevas_tarifas_sheets.append({
            'Ruta': ruta1_nombre,
            'Proveedor': p['prov'],
            'Punto_Sucursal': p['pto'],
            'Producto': item['prod_original'],
            'Precio_Kg': item['precio'],
            'Estado': item['estado'],
            'Observaciones': 'Tarifa heredada Santa Elena'
        })

# 4.2 Proveedores de Cavasa (8)
provs_cavasa = [
    {'prov': 'Barbara Gomez', 'pto': 'Cavasa'},
    {'prov': 'Los Lagos', 'pto': 'Cavasa'},
    {'prov': 'Caribe', 'pto': 'Cavasa'},
    {'prov': 'Sevillana', 'pto': 'Cavasa'},
    {'prov': 'Migan Capital', 'pto': 'Cavasa'},
    {'prov': 'Freddy Hernandez', 'pto': 'Cavasa'},
    {'prov': 'Edinson Aguirre', 'pto': 'Cavasa'},
    {'prov': 'La Reserva', 'pto': 'Cavasa'}
]

for p in provs_cavasa:
    for prod_n, item in sorted(tarifas_cavasa.items()):
        nuevas_tarifas_sheets.append({
            'Ruta': ruta1_nombre,
            'Proveedor': p['prov'],
            'Punto_Sucursal': p['pto'],
            'Producto': item['prod_original'],
            'Precio_Kg': item['precio'],
            'Estado': item['estado'],
            'Observaciones': 'Tarifa heredada Cavasa'
        })

# 4.3 Filas de respaldo histórico para Santa Elena y Cavasa
for prod_n, item in sorted(tarifas_santa_elena.items()):
    nuevas_tarifas_sheets.append({
        'Ruta': ruta1_nombre,
        'Proveedor': 'Santa Elena',
        'Punto_Sucursal': '',
        'Producto': item['prod_original'],
        'Precio_Kg': item['precio'],
        'Estado': item['estado'],
        'Observaciones': 'Tarifa respaldo histórica'
    })

for prod_n, item in sorted(tarifas_cavasa.items()):
    nuevas_tarifas_sheets.append({
        'Ruta': ruta1_nombre,
        'Proveedor': 'Cavasa',
        'Punto_Sucursal': '',
        'Producto': item['prod_original'],
        'Precio_Kg': item['precio'],
        'Estado': item['estado'],
        'Observaciones': 'Tarifa respaldo histórica'
    })

# 4.4 Agregar las demás rutas
for r in otras_tarifas_existentes:
    nuevas_tarifas_sheets.append({
        'Ruta': r['ruta'],
        'Proveedor': r['prov'],
        'Punto_Sucursal': r['pto'],
        'Producto': r['prod'],
        'Precio_Kg': r['precio'],
        'Estado': r['estado'],
        'Observaciones': r['obs']
    })

print(f"\n=== TOTAL FILAS MATRIZ TARIFAS DEPURADAS: {len(nuevas_tarifas_sheets)} ===")
print(f"  • Filas Ruta 1 (Nuevos 11 Proveedores): {len(provs_santa_elena)*len(tarifas_santa_elena) + len(provs_cavasa)*len(tarifas_cavasa)} filas")
print(f"  • Filas Ruta 1 (Respaldo Santa Elena + Cavasa): {len(tarifas_santa_elena) + len(tarifas_cavasa)} filas")
print(f"  • Filas Demás Rutas (Intactas): {len(otras_tarifas_existentes)} filas")

# ==============================================================================
# ESTILOS PROFESIONALES DE EXCEL
# ==============================================================================
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
fill_header = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
font_data = Font(name='Calibri', size=10)
font_bold = Font(name='Calibri', size=10, bold=True)
fill_pending = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
font_pending = Font(name='Calibri', size=10, color='92400E')
border_thin = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')
align_center = Alignment(horizontal='center', vertical='center')

def apply_sheet_styling(ws, headers, rows, is_provider_only=False):
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center if 'Precio' in headers[c_idx-1] or 'Estado' in headers[c_idx-1] else align_left
        cell.border = border_thin
        
    for r_idx, r_data in enumerate(rows, start=2):
        ws.row_dimensions[r_idx].height = 20
        precio_val = r_data.get('Precio_Kg')
        is_pending = (precio_val is None or precio_val == '')
        
        for c_idx, h in enumerate(headers, start=1):
            val = r_data.get(h, '')
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border_thin
            
            if h == 'Producto':
                cell.font = font_bold
                cell.alignment = align_left
            elif h == 'Precio_Kg':
                cell.font = font_bold if not is_pending else font_pending
                cell.alignment = align_right
                if not is_pending:
                    cell.number_format = '$#,##0'
            elif h == 'Estado':
                cell.font = font_pending if is_pending else font_data
                cell.alignment = align_center
                if is_pending:
                    cell.fill = fill_pending
            elif h == 'Observaciones':
                cell.font = font_pending if is_pending else font_data
                cell.alignment = align_left
                if is_pending:
                    cell.fill = fill_pending
            else:
                cell.font = font_data
                cell.alignment = align_left

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# ==============================================================================
# 5. ACTUALIZAR Tarifario_Completo_Reparado_2026.xlsx
# ==============================================================================
print("\n=== 5. GENERANDO Tarifario_Completo_Reparado_2026.xlsx ===")
# Generar producto cartesiano: 126 puntos x 22 materias = 2,772 filas
filas_cartesiano_puntos = []

# Mapas de tarifas para búsqueda exacta en el cartesiano
mapa_tarifas_exactas = {}
for item in nuevas_tarifas_sheets:
    k = (norm(item['Proveedor']), norm(item['Punto_Sucursal']), norm(item['Producto']))
    if item['Precio_Kg'] is not None and item['Precio_Kg'] != '':
        mapa_tarifas_exactas[k] = item

for pt in puntos_oficiales:
    ruta = pt['ruta']
    prov = pt['prov']
    pto = pt['pto']
    
    for prod in prods_activos:
        prod_n = norm(prod)
        prov_n = norm(prov)
        pto_n = norm(pto)
        
        # Búsqueda exacta
        precio = None
        estado = 'Pendiente Precio'
        obs = 'Pendiente cotizar'
        
        k_ex = (prov_n, pto_n, prod_n)
        if k_ex in mapa_tarifas_exactas:
            precio = mapa_tarifas_exactas[k_ex]['Precio_Kg']
            estado = 'Activo'
            obs = mapa_tarifas_exactas[k_ex]['Observaciones']
        elif (prov_n, '', prod_n) in mapa_tarifas_exactas:
            precio = mapa_tarifas_exactas[(prov_n, '', prod_n)]['Precio_Kg']
            estado = 'Activo'
            obs = 'Tarifa proveedor general'
        else:
            # Fallback por sector Santa Elena / Cavasa
            if 'Ruta 1' in ruta:
                if prov_n in ['BODEGA SANTA ELENA', 'ALEJANDRO GARAY', 'SEVILLANA SANTA ELENA'] or pto_n == 'SANTA ELENA':
                    if prod_n in tarifas_santa_elena and tarifas_santa_elena[prod_n]['precio']:
                        precio = tarifas_santa_elena[prod_n]['precio']
                        estado = 'Activo'
                        obs = 'Tarifa heredada Santa Elena'
                elif pto_n == 'CAVASA' or prov_n in ['BARBARA GOMEZ', 'LOS LAGOS', 'CARIBE', 'SEVILLANA', 'MIGAN CAPITAL', 'FREDDY HERNANDEZ', 'EDINSON AGUIRRE', 'LA RESERVA']:
                    if prod_n in tarifas_cavasa and tarifas_cavasa[prod_n]['precio']:
                        precio = tarifas_cavasa[prod_n]['precio']
                        estado = 'Activo'
                        obs = 'Tarifa heredada Cavasa'
        
        filas_cartesiano_puntos.append({
            'Ruta': ruta,
            'Proveedor': prov,
            'Punto_Sucursal': pto,
            'Producto': prod,
            'Precio_Kg': precio,
            'Estado': estado,
            'Observaciones': obs
        })

wb_comp = openpyxl.Workbook()
# Tab 1: Tarifas (Cartesiano completo 2,772 filas)
ws_tc = wb_comp.active
ws_tc.title = 'Tarifas'
apply_sheet_styling(ws_tc, ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'], filas_cartesiano_puntos)

# Tab 2: Tarifas_Depuradas (Matriz limpia 372 filas)
ws_td = wb_comp.create_sheet('Tarifas_Depuradas')
apply_sheet_styling(ws_td, ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'], nuevas_tarifas_sheets)

# Tab 3: Puntos_Rutas
ws_pr = wb_comp.create_sheet('Puntos_Rutas')
headers_pr = ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Direccion', 'Telefono', 'Horario_Estimado', 'Frecuencia_Dias', 'Estado', 'Materias_Frecuentes']
ws_pr.append(headers_pr)
ws_pr.row_dimensions[1].height = 26
for c_idx in range(1, len(headers_pr) + 1):
    c = ws_pr.cell(row=1, column=c_idx)
    c.font = font_header
    c.fill = fill_header
    c.border = border_thin
for r_idx, pt in enumerate(puntos_oficiales, start=2):
    ws_pr.row_dimensions[r_idx].height = 20
    vals = [pt['ruta'], pt['prov'], pt['pto'], pt['dir'], pt['tel'], pt['hora'], pt['frec'], pt['estado'], pt['materias']]
    for c_idx, val in enumerate(vals, start=1):
        cell = ws_pr.cell(row=r_idx, column=c_idx, value=val)
        cell.font = font_data
        cell.border = border_thin

# Tab 4: Productos
ws_p = wb_comp.create_sheet('Productos')
ws_p.append(['Producto', 'Estado'])
ws_p.row_dimensions[1].height = 26
for c_idx in [1, 2]:
    c = ws_p.cell(row=1, column=c_idx)
    c.font = font_header
    c.fill = fill_header
    c.border = border_thin
for r_idx, r in enumerate(prods_raw[1:], start=2):
    ws_p.row_dimensions[r_idx].height = 20
    c1 = ws_p.cell(row=r_idx, column=1, value=r[0] if len(r)>0 else '')
    c2 = ws_p.cell(row=r_idx, column=2, value=r[1] if len(r)>1 else '')
    c1.font = font_bold; c1.border = border_thin
    c2.font = font_data; c2.border = border_thin

try:
    wb_comp.save('Tarifario_Completo_Reparado_2026.xlsx')
    print("✅ Guardado 'Tarifario_Completo_Reparado_2026.xlsx'")
except Exception as e:
    wb_comp.save('Tarifario_Completo_Reparado_2026_Actualizado.xlsx')
    print("⚠️ 'Tarifario_Completo_Reparado_2026.xlsx' está abierto en Excel. Guardado como 'Tarifario_Completo_Reparado_2026_Actualizado.xlsx'")

# ==============================================================================
# 6. ACTUALIZAR Tarifario_Por_Proveedor_2026.xlsx
# ==============================================================================
print("\n=== 6. GENERANDO Tarifario_Por_Proveedor_2026.xlsx ===")
prov_dict = {}
for pt in puntos_oficiales:
    prov = pt['prov']
    ruta = pt['ruta']
    pto = pt['pto']
    if prov not in prov_dict:
        prov_dict[prov] = {'rutas': set(), 'puntos': set()}
    prov_dict[prov]['rutas'].add(ruta)
    prov_dict[prov]['puntos'].add(pto)

for prov, d in prov_dict.items():
    if len(d['rutas']) == 1:
        d['ruta_display'] = list(d['rutas'])[0]
    else:
        d['ruta_display'] = "Multirruta (" + ", ".join(sorted([r.split(':')[0] for r in d['rutas']])) + ")"

filas_prov_clean = []
filas_prov_general = []

for prov in sorted(prov_dict.keys()):
    d = prov_dict[prov]
    ruta_str = d['ruta_display']
    prov_n = norm(prov)
    
    for prod in prods_activos:
        prod_n = norm(prod)
        precio = None
        estado = 'Pendiente Precio'
        obs = 'Pendiente cotizar'
        
        # Búsqueda en mapa
        # Directa por proveedor
        encontrado = False
        for (pr_k, pto_k, pd_k), item in mapa_tarifas_exactas.items():
            if pr_k == prov_n and pd_k == prod_n:
                precio = item['Precio_Kg']
                estado = 'Activo'
                obs = item['Observaciones']
                encontrado = True
                break
                
        if not encontrado:
            if prov_n in ['BODEGA SANTA ELENA', 'ALEJANDRO GARAY', 'SEVILLANA SANTA ELENA']:
                if prod_n in tarifas_santa_elena and tarifas_santa_elena[prod_n]['precio']:
                    precio = tarifas_santa_elena[prod_n]['precio']
                    estado = 'Activo'
                    obs = 'Tarifa heredada Santa Elena'
            elif prov_n in ['BARBARA GOMEZ', 'LOS LAGOS', 'CARIBE', 'SEVILLANA', 'MIGAN CAPITAL', 'FREDDY HERNANDEZ', 'EDINSON AGUIRRE', 'LA RESERVA'] or 'CAVASA' in [norm(pt) for pt in d['puntos']]:
                if prod_n in tarifas_cavasa and tarifas_cavasa[prod_n]['precio']:
                    precio = tarifas_cavasa[prod_n]['precio']
                    estado = 'Activo'
                    obs = 'Tarifa heredada Cavasa'
                    
        filas_prov_clean.append({
            'Ruta': ruta_str,
            'Proveedor': prov,
            'Producto': prod,
            'Precio_Kg': precio,
            'Estado': estado,
            'Observaciones': obs
        })
        
        filas_prov_general.append({
            'Ruta': ruta_str,
            'Proveedor': prov,
            'Punto_Sucursal': 'General',
            'Producto': prod,
            'Precio_Kg': precio,
            'Estado': estado,
            'Observaciones': obs
        })

wb_prov = openpyxl.Workbook()
# Tab 1: Tarifas_Por_Proveedor (6 columnas)
ws_pp = wb_prov.active
ws_pp.title = 'Tarifas_Por_Proveedor'
apply_sheet_styling(ws_pp, ['Ruta', 'Proveedor', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'], filas_prov_clean)

# Tab 2: Tarifas_Punto_General (7 columnas)
ws_pg = wb_prov.create_sheet('Tarifas_Punto_General')
apply_sheet_styling(ws_pg, ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'], filas_prov_general)

# Tab 3: Tarifas_Depuradas
ws_p_dep = wb_prov.create_sheet('Tarifas_Depuradas')
apply_sheet_styling(ws_p_dep, ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'], nuevas_tarifas_sheets)

# Tab 4: Puntos_Rutas
ws_p_pr = wb_prov.create_sheet('Puntos_Rutas')
ws_p_pr.append(headers_pr)
ws_p_pr.row_dimensions[1].height = 26
for c_idx in range(1, len(headers_pr) + 1):
    c = ws_p_pr.cell(row=1, column=c_idx)
    c.font = font_header; c.fill = fill_header; c.border = border_thin
for r_idx, pt in enumerate(puntos_oficiales, start=2):
    ws_p_pr.row_dimensions[r_idx].height = 20
    vals = [pt['ruta'], pt['prov'], pt['pto'], pt['dir'], pt['tel'], pt['hora'], pt['frec'], pt['estado'], pt['materias']]
    for c_idx, val in enumerate(vals, start=1):
        cell = ws_p_pr.cell(row=r_idx, column=c_idx, value=val)
        cell.font = font_data; cell.border = border_thin

# Tab 5: Productos
ws_p_prod = wb_prov.create_sheet('Productos')
ws_p_prod.append(['Producto', 'Estado'])
ws_p_prod.row_dimensions[1].height = 26
for c_idx in [1, 2]:
    c = ws_p_prod.cell(row=1, column=c_idx)
    c.font = font_header; c.fill = fill_header; c.border = border_thin
for r_idx, r in enumerate(prods_raw[1:], start=2):
    ws_p_prod.row_dimensions[r_idx].height = 20
    c1 = ws_p_prod.cell(row=r_idx, column=1, value=r[0] if len(r)>0 else '')
    c2 = ws_p_prod.cell(row=r_idx, column=2, value=r[1] if len(r)>1 else '')
    c1.font = font_bold; c1.border = border_thin
    c2.font = font_data; c2.border = border_thin

try:
    wb_prov.save('Tarifario_Por_Proveedor_2026.xlsx')
    print("✅ Guardado 'Tarifario_Por_Proveedor_2026.xlsx'")
except Exception as e:
    wb_prov.save('Tarifario_Por_Proveedor_2026_Actualizado.xlsx')
    print("⚠️ 'Tarifario_Por_Proveedor_2026.xlsx' está abierto en Excel. Guardado como 'Tarifario_Por_Proveedor_2026_Actualizado.xlsx'")

# ==============================================================================
# 7. GENERAR CSV Y EXCEL PARA IMPORTACIÓN DIRECTA A GOOGLE SHEETS
# ==============================================================================
wb_gs = openpyxl.Workbook()
ws_gs = wb_gs.active
ws_gs.title = 'Tarifas'
apply_sheet_styling(ws_gs, ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'], nuevas_tarifas_sheets)
try:
    wb_gs.save('Tarifas_GoogleSheets_Actualizado.xlsx')
    print("✅ Guardado 'Tarifas_GoogleSheets_Actualizado.xlsx'")
except Exception as e:
    wb_gs.save('Tarifas_GoogleSheets_Actualizado_v2.xlsx')
    print("⚠️ Guardado como 'Tarifas_GoogleSheets_Actualizado_v2.xlsx'")

try:
    with open('Tarifas_GoogleSheets_Actualizado.csv', 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'])
        for r in nuevas_tarifas_sheets:
            p_str = f"${r['Precio_Kg']:,}".replace(',', '.') if r['Precio_Kg'] is not None else ''
            writer.writerow([r['Ruta'], r['Proveedor'], r['Punto_Sucursal'], r['Producto'], p_str, r['Estado'], r['Observaciones']])
    print("✅ Generado 'Tarifas_GoogleSheets_Actualizado.csv'")
except Exception as e:
    print("⚠️ Error guardando CSV:", e)

# Actualizar también tarifario.xlsx con la nueva hoja Tarifas y Puntos_Rutas
try:
    wb_base = openpyxl.load_workbook('tarifario.xlsx')
    if 'Tarifas' in wb_base.sheetnames:
        del wb_base['Tarifas']
    ws_base_t = wb_base.create_sheet('Tarifas', 0)
    apply_sheet_styling(ws_base_t, ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones'], nuevas_tarifas_sheets)
    
    if 'Puntos_Rutas' in wb_base.sheetnames:
        del wb_base['Puntos_Rutas']
    ws_base_pr = wb_base.create_sheet('Puntos_Rutas', 2)
    ws_base_pr.append(headers_pr)
    ws_base_pr.row_dimensions[1].height = 26
    for c_idx in range(1, len(headers_pr) + 1):
        c = ws_base_pr.cell(row=1, column=c_idx)
        c.font = font_header; c.fill = fill_header; c.border = border_thin
    for r_idx, pt in enumerate(puntos_oficiales, start=2):
        ws_base_pr.row_dimensions[r_idx].height = 20
        vals = [pt['ruta'], pt['prov'], pt['pto'], pt['dir'], pt['tel'], pt['hora'], pt['frec'], pt['estado'], pt['materias']]
        for c_idx, val in enumerate(vals, start=1):
            cell = ws_base_pr.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_data; cell.border = border_thin
            
    wb_base.save('tarifario.xlsx')
    print("✅ Guardado master 'tarifario.xlsx' con Tarifas y Puntos_Rutas actualizados.")
except Exception as e:
    print("⚠️ Error actualizando tarifario.xlsx:", e)

print("\n🎉 ¡PROCESO DE ACTUALIZACIÓN DE TARIFAS RUTA 1 COMPLETADO CON ÉXITO!")
