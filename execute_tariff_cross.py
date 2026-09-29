import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import unicodedata
import re

def clean_text(s):
    if s is None:
        return ''
    s = str(s).strip()
    s = s.replace('hist\ufffdrico', 'histórico').replace('Hist\ufffdrico', 'Histórico')
    s = s.replace('ca\ufffda', 'caña').replace('Ca\ufffda', 'Caña').replace('CA\ufffdA', 'CAÑA')
    s = s.replace('le\ufffda', 'leña').replace('Le\ufffda', 'Leña').replace('LE\ufffdA', 'LEÑA')
    s = s.replace('s\ufffdbado', 'sábado').replace('S\ufffdbado', 'Sábado')
    s = s.replace('mill\ufffdn', 'millán').replace('Mill\ufffdn', 'Millán')
    s = s.replace('jamund\ufffd', 'jamundí').replace('Jamund\ufffd', 'Jamundí')
    s = s.replace('guacar\ufffd', 'guacarí').replace('Guacar\ufffd', 'Guacarí')
    s = s.replace('r\ufffdo', 'río').replace('R\ufffdo', 'Río')
    s = s.replace('berl\ufffdn', 'berlín').replace('Berl\ufffdn', 'Berlín')
    s = s.replace('tulu\ufffd', 'tuluá').replace('Tulu\ufffd', 'Tuluá')
    s = s.replace('jes\ufffds', 'jesús').replace('Jes\ufffds', 'Jesús')
    s = s.replace('jos\ufffd', 'josé').replace('Jos\ufffd', 'José')
    s = s.replace('rep\ufffdblica', 'república').replace('Rep\ufffdblica', 'República')
    s = s.replace('rebaj\ufffdn', 'rebajón').replace('Rebaj\ufffdn', 'Rebajón')
    s = s.replace('fabi\ufffdn', 'fabián').replace('Fabi\ufffdn', 'Fabián')
    s = s.replace('l\ufffdpez', 'lópez').replace('L\ufffdpez', 'López')
    s = s.replace('adri\ufffdn', 'adrián').replace('Adri\ufffdn', 'Adrián')
    s = s.replace('uni\ufffdn', 'unión').replace('Uni\ufffdn', 'Unión')
    s = s.replace('mu\ufffdoz', 'muñoz').replace('Mu\ufffdoz', 'Muñoz')
    s = s.replace('y\ufffdnifer', 'yénifer').replace('Y\ufffdnifer', 'Yénifer')
    s = s.replace('d\ufffdaz', 'díaz').replace('D\ufffdaz', 'Díaz')
    s = s.replace('frigor\uffdfico', 'frigorífico').replace('Frigor\uffdfico', 'Frigorífico')
    s = s.replace('galer\ufffda', 'galería').replace('Galer\ufffda', 'Galería')
    s = s.replace('bol\ufffdvar', 'bolívar').replace('Bol\ufffdvar', 'Bolívar')
    s = s.replace('guacan\ufffda', 'guacandá').replace('Guacan\ufffda', 'Guacandá')
    s = s.replace('monta\ufffda', 'montaña').replace('Monta\ufffda', 'Montaña')
    s = s.replace('\ufffd', '')
    return s

def norm(s):
    if not s: return ''
    s = clean_text(s)
    nfkd = unicodedata.normalize('NFKD', s)
    res = ''.join([c for c in nfkd if not unicodedata.combining(c)])
    return ' '.join(re.sub(r'[^a-zA-Z0-9]', ' ', res).upper().split())

wb = openpyxl.load_workbook('tarifario.xlsx', data_only=True)
ws_t = wb['Tarifas']
ws_p = wb['Puntos_Rutas']
ws_prod = wb['Productos']

# 1. Active Products
prods_activos = []
for r in range(2, ws_prod.max_row + 1):
    val = ws_prod.cell(row=r, column=1).value
    st = ws_prod.cell(row=r, column=2).value
    if val and st and clean_text(st).lower() == 'activo':
        name = clean_text(val).upper()
        if 'LE' in name and 'A' in name and len(name) <= 5:
            name = 'LEÑA'
        prods_activos.append(name)
prods_activos = sorted(list(set(prods_activos)))

# 2. Extract strictly UNIQUE 62 Providers from Puntos_Rutas
prov_data = {}
for r in range(2, ws_p.max_row + 1):
    ruta = clean_text(ws_p.cell(row=r, column=1).value)
    prov = clean_text(ws_p.cell(row=r, column=2).value)
    pto = clean_text(ws_p.cell(row=r, column=3).value)
    if prov:
        if prov not in prov_data:
            prov_data[prov] = {'rutas': set(), 'puntos': set()}
        if ruta: prov_data[prov]['rutas'].add(ruta)
        if pto: prov_data[prov]['puntos'].add(pto)

for prov, data in prov_data.items():
    rutas_list = sorted(list(data['rutas']))
    if len(rutas_list) == 1:
        data['ruta_display'] = rutas_list[0]
    elif len(rutas_list) > 1:
        data['ruta_display'] = "Multirruta (" + ", ".join([r.split(':')[0] for r in rutas_list]) + ")"
    else:
        data['ruta_display'] = 'Sin ruta'

# 3. Read every single price from source sheet Tarifas
tarifas_extracted = []
for r in range(2, ws_t.max_row + 1):
    ruta = clean_text(ws_t.cell(row=r, column=1).value)
    prov = clean_text(ws_t.cell(row=r, column=2).value)
    pto = clean_text(ws_t.cell(row=r, column=3).value)
    prod = clean_text(ws_t.cell(row=r, column=4).value).upper()
    precio = ws_t.cell(row=r, column=5).value
    estado = clean_text(ws_t.cell(row=r, column=6).value)
    obs = clean_text(ws_t.cell(row=r, column=7).value)
    
    if precio not in [None, '', 0, '0', 'Pendiente Precio'] and prov and prod:
        try:
            p_num = float(precio)
            if p_num.is_integer(): p_num = int(p_num)
        except:
            p_num = precio
            
        tarifas_extracted.append({
            'row': r,
            'ruta': ruta,
            'prov': prov,
            'pto': pto,
            'prod': prod,
            'precio': p_num,
            'estado': estado,
            'obs': obs
        })

print(f"Total registros con precio en hoja origen 'Tarifas': {len(tarifas_extracted)}")

def normalize_prod(p):
    pn = norm(p)
    if pn in ['HUESO', 'HUESO ']:
        return 'HUESO BLANCO'
    if 'ACEITE' in pn:
        return 'ACEITE'
    return pn

by_prov_prod = {}
by_pto_prod = {}

for t in tarifas_extracted:
    p_norm = normalize_prod(t['prod'])
    pr_norm = norm(t['prov'])
    pt_norm = norm(t['pto'])
    
    by_prov_prod.setdefault((pr_norm, p_norm), []).append(t)
    if pt_norm:
        by_pto_prod.setdefault((pt_norm, p_norm), []).append(t)

# 4. Perform Complete Cross-Matching (Cruce)
rows_clean = []
crossed_count = 0
pending_count = 0

for prov in sorted(prov_data.keys()):
    pr_norm = norm(prov)
    ptos = prov_data[prov]['puntos']
    ruta_str = prov_data[prov]['ruta_display']
    
    for prod in prods_activos:
        p_norm = normalize_prod(prod)
        found_price = None
        origen_desc = ''
        
        # 1. Direct match (Proveedor + Producto)
        if (pr_norm, p_norm) in by_prov_prod:
            t = by_prov_prod[(pr_norm, p_norm)][0]
            found_price = t['precio']
            origen_desc = 'Tarifa cruzada de hoja Tarifas'
            
        # 2. Check points/sucursales of this provider
        if found_price is None:
            for pt in ptos:
                pt_n = norm(pt)
                if (pt_n, p_norm) in by_pto_prod:
                    t = by_pto_prod[(pt_n, p_norm)][0]
                    found_price = t['precio']
                    origen_desc = 'Tarifa cruzada de hoja Tarifas'
                    break
                    
        # 3. Known Administrative Aliases in sheet Tarifas
        if found_price is None:
            if 'SEVILLANA' in pr_norm and (norm('Cuenta Sevillana'), p_norm) in by_prov_prod:
                t = by_prov_prod[(norm('Cuenta Sevillana'), p_norm)][0]
                found_price = t['precio']
                origen_desc = 'Tarifa cruzada de hoja Tarifas'
            elif 'ALBERTO MILLAN' in pr_norm and (norm('Cuenta Alberto Millan'), p_norm) in by_prov_prod:
                t = by_prov_prod[(norm('Cuenta Alberto Millan'), p_norm)][0]
                found_price = t['precio']
                origen_desc = 'Tarifa cruzada de hoja Tarifas'
            elif 'BELALCAZAR' in pr_norm and (norm('Belalcazar'), p_norm) in by_prov_prod:
                t = by_prov_prod[(norm('Belalcazar'), p_norm)][0]
                found_price = t['precio']
                origen_desc = 'Tarifa cruzada de hoja Tarifas'
            elif 'FRIGORIVALLE' in pr_norm and (norm('Supertienda Cañaveral'), p_norm) in by_prov_prod:
                t = by_prov_prod[(norm('Supertienda Cañaveral'), p_norm)][0]
                found_price = t['precio']
                origen_desc = 'Tarifa cruzada de hoja Tarifas'

        if found_price is not None:
            crossed_count += 1
            rows_clean.append({
                'Ruta': ruta_str,
                'Proveedor': prov,
                'Producto': prod,
                'Precio_Kg': found_price,
                'Estado': 'Activo',
                'Observaciones': origen_desc
            })
        else:
            pending_count += 1
            rows_clean.append({
                'Ruta': ruta_str,
                'Proveedor': prov,
                'Producto': prod,
                'Precio_Kg': None,
                'Estado': 'Pendiente Precio',
                'Observaciones': 'Pendiente cotizar'
            })

print(f"Total combinaciones generadas: {len(rows_clean)} (exactamente {len(prov_data)} proveedores × {len(prods_activos)} productos)")
print(f"Tarifas exitosamente cruzadas con precio: {crossed_count}")
print(f"Tarifas pendientes: {pending_count}")

# 5. Styling
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
fill_header = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
font_data = Font(name='Calibri', size=10)
font_bold = Font(name='Calibri', size=10, bold=True)
fill_pending = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
font_pending = Font(name='Calibri', size=10, color='92400E')
border_thin = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')
align_center = Alignment(horizontal='center', vertical='center')

def populate_sheet(ws, rows_data):
    ws.delete_rows(1, ws.max_row + 10)
    headers = ['Ruta', 'Proveedor', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones']
    ws.append(headers)
    
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center if headers[c_idx-1] in ['Precio_Kg', 'Estado'] else align_left
        cell.border = border_thin
    ws.row_dimensions[1].height = 26
    
    for r_idx, r in enumerate(rows_data, start=2):
        ws.row_dimensions[r_idx].height = 20
        is_p = (r['Precio_Kg'] is None)
        
        c1 = ws.cell(row=r_idx, column=1, value=r['Ruta'])
        c1.font = font_data; c1.alignment = align_left; c1.border = border_thin
        
        c2 = ws.cell(row=r_idx, column=2, value=r['Proveedor'])
        c2.font = font_data; c2.alignment = align_left; c2.border = border_thin
        
        c3 = ws.cell(row=r_idx, column=3, value=r['Producto'])
        c3.font = font_bold; c3.alignment = align_left; c3.border = border_thin
        
        c4 = ws.cell(row=r_idx, column=4, value=r['Precio_Kg'])
        c4.font = font_bold if not is_p else font_pending; c4.alignment = align_right; c4.border = border_thin
        if not is_p: c4.number_format = '$#,##0'
        
        c5 = ws.cell(row=r_idx, column=5, value=r['Estado'])
        c5.font = font_pending if is_p else font_data; c5.alignment = align_center; c5.border = border_thin
        if is_p: c5.fill = fill_pending
        
        c6 = ws.cell(row=r_idx, column=6, value=r['Observaciones'])
        c6.font = font_pending if is_p else font_data; c6.alignment = align_left; c6.border = border_thin

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:F{len(rows_data)+1}"
    
    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 32
    ws.column_dimensions['C'].width = 26
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 34

# 6. Apply to tarifario.xlsx -> Tarifas_Por_Proveedor
if 'Tarifas_Por_Proveedor' in wb.sheetnames:
    ws_tp = wb['Tarifas_Por_Proveedor']
else:
    ws_tp = wb.create_sheet('Tarifas_Por_Proveedor')

populate_sheet(ws_tp, rows_clean)

# Also ensure Tarifas_Punto_General is clean and updated
rows_gen = []
for r in rows_clean:
    r_g = dict(r)
    r_g['Punto_Sucursal'] = 'General'
    rows_gen.append(r_g)

if 'Tarifas_Punto_General' in wb.sheetnames:
    ws_gen = wb['Tarifas_Punto_General']
    ws_gen.delete_rows(1, ws_gen.max_row + 10)
    g_headers = ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones']
    ws_gen.append(g_headers)
    for c_idx in range(1, len(g_headers) + 1):
        cell = ws_gen.cell(row=1, column=c_idx)
        cell.font = font_header; cell.fill = fill_header; cell.border = border_thin
        cell.alignment = align_center if g_headers[c_idx-1] in ['Precio_Kg', 'Estado'] else align_left
    ws_gen.row_dimensions[1].height = 26
    
    for r_idx, r in enumerate(rows_gen, start=2):
        ws_gen.row_dimensions[r_idx].height = 20
        is_p = (r['Precio_Kg'] is None)
        for c_idx, h in enumerate(g_headers, start=1):
            val = r.get(h)
            cell = ws_gen.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border_thin
            if h == 'Precio_Kg':
                cell.font = font_bold if not is_p else font_pending
                cell.alignment = align_right
                if not is_p: cell.number_format = '$#,##0'
            elif h == 'Estado':
                cell.font = font_pending if is_p else font_data
                cell.alignment = align_center
                if is_p: cell.fill = fill_pending
            elif h == 'Producto':
                cell.font = font_bold; cell.alignment = align_left
            elif h == 'Observaciones':
                cell.font = font_pending if is_p else font_data; cell.alignment = align_left
            else:
                cell.font = font_data; cell.alignment = align_left
                
    ws_gen.freeze_panes = 'A2'
    ws_gen.auto_filter.ref = f"A1:G{len(rows_gen)+1}"
    ws_gen.column_dimensions['A'].width = 38
    ws_gen.column_dimensions['B'].width = 30
    ws_gen.column_dimensions['C'].width = 18
    ws_gen.column_dimensions['D'].width = 26
    ws_gen.column_dimensions['E'].width = 16
    ws_gen.column_dimensions['F'].width = 18
    ws_gen.column_dimensions['G'].width = 34

wb.save('tarifario.xlsx')
print("Archivo tarifario.xlsx actualizado exitosamente.")

# Also update Tarifario_Por_Proveedor_2026.xlsx and Tarifario_Completo_Reparado_2026.xlsx
for f_out in ['Tarifario_Por_Proveedor_2026.xlsx', 'Tarifario_Completo_Reparado_2026.xlsx']:
    try:
        wb_copy = openpyxl.load_workbook(f_out)
        if 'Tarifas_Por_Proveedor' in wb_copy.sheetnames:
            populate_sheet(wb_copy['Tarifas_Por_Proveedor'], rows_clean)
        wb_copy.save(f_out)
        print(f"Copia actualizada: {f_out}")
    except Exception as e:
        print(f"No se pudo actualizar {f_out}: {e}")
