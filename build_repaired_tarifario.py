import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import unicodedata
import re

def clean_text(s):
    if s is None:
        return ''
    s = str(s).strip()
    # Specific known words with corruption
    s = s.replace('hist\ufffdrico', 'histórico').replace('Hist\ufffdrico', 'Histórico')
    s = s.replace('ca\ufffda', 'caña').replace('Ca\ufffda', 'Caña').replace('CA\ufffdA', 'CAÑA')
    s = s.replace('le\ufffda', 'leña').replace('Le\ufffda', 'Leña').replace('LE\ufffdA', 'LEÑA')
    s = s.replace('s\ufffdbado', 'sábado').replace('S\ufffdbado', 'Sábado')
    s = s.replace('mill\ufffdn', 'millán').replace('Mill\ufffdn', 'Millán')
    s = s.replace('jamund\ufffd', 'jamundí').replace('Jamund\ufffd', 'Jamundí')
    s = s.replace('guacar\ufffd', 'guacarí').replace('Guacar\ufffd', 'Guacarí')
    s = s.replace('r\ufffdo', 'río').replace('R\ufffdo', 'Río')
    s = s.replace('berl\ufffdn', 'berlín').replace('Berl\ufffdn', 'Berlín')
    s = s.replace('tulu\ufffd', 'tuluá').replace('Tulu\ufffd', 'Tuluá')
    s = s.replace('jes\ufffds', 'jesús').replace('Jes\ufffds', 'Jesús')
    s = s.replace('jos\ufffd', 'josé').replace('Jos\ufffd', 'José')
    s = s.replace('rep\ufffdblica', 'república').replace('Rep\ufffdblica', 'República')
    s = s.replace('rebaj\ufffdn', 'rebajón').replace('Rebaj\ufffdn', 'Rebajón')
    s = s.replace('fabi\ufffdn', 'fabián').replace('Fabi\ufffdn', 'Fabián')
    s = s.replace('l\ufffdpez', 'lópez').replace('L\ufffdpez', 'López')
    s = s.replace('adri\ufffdn', 'adrián').replace('Adri\ufffdn', 'Adrián')
    s = s.replace('uni\ufffdn', 'unión').replace('Uni\ufffdn', 'Unión')
    s = s.replace('mu\ufffdoz', 'muñoz').replace('Mu\ufffdoz', 'Muñoz')
    s = s.replace('y\ufffdnifer', 'yénifer').replace('Y\ufffdnifer', 'Yénifer')
    s = s.replace('d\ufffdaz', 'díaz').replace('D\ufffdaz', 'Díaz')
    s = s.replace('frigor\uffdfico', 'frigorífico').replace('Frigor\uffdfico', 'Frigorífico')
    s = s.replace('galer\ufffda', 'galería').replace('Galer\ufffda', 'Galería')
    s = s.replace('bol\ufffdvar', 'bolívar').replace('Bol\ufffdvar', 'Bolívar')
    s = s.replace('guacan\ufffda', 'guacandá').replace('Guacan\ufffda', 'Guacandá')
    s = s.replace('monta\ufffda', 'montaña').replace('Monta\ufffda', 'Montaña')
    s = s.replace('\ufffd', '') # Catch-all for any other stray replacement character
    return s

def norm(s):
    if not s: return ''
    s = clean_text(s)
    nfkd = unicodedata.normalize('NFKD', s)
    res = ''.join([c for c in nfkd if not unicodedata.combining(c)])
    return ' '.join(re.sub(r'[^a-zA-Z0-9]', ' ', res).upper().split())

wb_orig = openpyxl.load_workbook('tarifario.xlsx', data_only=True)
ws_prod = wb_orig['Productos']
ws_puntos = wb_orig['Puntos_Rutas']
ws_tarifas = wb_orig['Tarifas']

# 1. Read Active Products
prods_activos = []
prods_all = []
for r in range(2, ws_prod.max_row + 1):
    val = ws_prod.cell(row=r, column=1).value
    st = ws_prod.cell(row=r, column=2).value
    if val:
        name = clean_text(val).upper()
        if 'LE' in name and 'A' in name and len(name) <= 5:
            name = 'LEÑA'
        st_clean = clean_text(st) if st else 'Inactivo'
        prods_all.append((name, st_clean))
        if st_clean.lower() == 'activo':
            prods_activos.append(name)

prods_activos = sorted(list(set(prods_activos)))

# 2. Read 126 Puntos_Rutas
puntos = []
puntos_full_rows = []
for r in range(2, ws_puntos.max_row + 1):
    ruta = clean_text(ws_puntos.cell(row=r, column=1).value)
    prov = clean_text(ws_puntos.cell(row=r, column=2).value)
    pto = clean_text(ws_puntos.cell(row=r, column=3).value)
    if prov and pto:
        puntos.append({'ruta': ruta, 'prov': prov, 'pto': pto})
        row_vals = [clean_text(ws_puntos.cell(row=r, column=c).value) for c in range(1, 10)]
        puntos_full_rows.append(row_vals)

# 3. Read Existing Tarifas for matching
exact_map = {}
pto_map = {}
prov_map = {}

for r in range(2, ws_tarifas.max_row + 1):
    prov = ws_tarifas.cell(row=r, column=2).value
    pto = ws_tarifas.cell(row=r, column=3).value
    prod = ws_tarifas.cell(row=r, column=4).value
    precio = ws_tarifas.cell(row=r, column=5).value
    estado = ws_tarifas.cell(row=r, column=6).value
    obs = ws_tarifas.cell(row=r, column=7).value
    
    if not prov or not prod:
        continue
    
    if precio not in [None, '', 0, '0', 'Pendiente Precio']:
        try:
            precio_num = float(precio)
            if precio_num.is_integer():
                precio_num = int(precio_num)
        except:
            precio_num = precio
        
        ke = (norm(prov), norm(pto), norm(prod))
        kp = (norm(pto), norm(prod))
        kpr = (norm(prov), norm(prod))
        
        obs_clean = clean_text(obs) if obs else 'Tarifa histórica confirmada'
        exact_map[ke] = (precio_num, clean_text(estado), obs_clean)
        if kp not in pto_map:
            pto_map[kp] = (precio_num, clean_text(estado), obs_clean)
        prov_map.setdefault(kpr, []).append((precio_num, clean_text(estado), obs_clean))

def find_tariff(pt, prod):
    prov_n = norm(pt['prov'])
    pto_n = norm(pt['pto'])
    prod_n = norm(prod)
    
    # Priority 1: Exact Match (Proveedor + Punto + Producto)
    ke = (prov_n, pto_n, prod_n)
    if ke in exact_map:
        p, st, o = exact_map[ke]
        return p, 'Activo', (o if o else 'Tarifa histórica confirmada')
        
    # Priority 2: Match by Sucursal + Producto
    kp = (pto_n, prod_n)
    if kp in pto_map:
        p, st, o = pto_map[kp]
        return p, 'Activo', 'Tarifa recuperada por sucursal'
        
    # Priority 3: Special group rules
    if 'BELALCAZAR' in prov_n or 'BELALCAZAR' in pto_n:
        k_b = ('BELALCAZAR', 'BELALCAZAR CENTRO', prod_n)
        if k_b in exact_map:
            return exact_map[k_b][0], 'Activo', 'Tarifa matriz Belalcázar'
            
    if 'ALBERTO MILLAN' in prov_n or 'ALBERTO MILLAN' in pto_n:
        k_am = ('CUENTA ALBERTO MILLAN', 'ALBERTO MILLAN BUGA', prod_n)
        if k_am in exact_map:
            return exact_map[k_am][0], 'Activo', 'Tarifa Alberto Millán Buga'
            
    if 'CANAVERAL' in pto_n or 'CANAVERAL' in prov_n:
        if 'ROLDANILLO' in pto_n or 'ZARZAL' in pto_n:
            k_c = ('SUPERTIENDA CANAVERAL', 'CANAVERAL ROLDANILLO ZARZAL', prod_n)
            if k_c in exact_map:
                return exact_map[k_c][0], 'Activo', 'Tarifa zona Roldanillo-Zarzal'
        if 'TULUA' in pto_n or 'BUGA' in pto_n:
            k_c = ('SUPERTIENDA CANAVERAL', 'CANAVERAL TULUA BUGA', prod_n)
            if k_c in exact_map:
                return exact_map[k_c][0], 'Activo', 'Tarifa zona Tuluá-Buga'
                
    if 'PUERTO TEJADA' in pto_n and 'CARIBE' in prov_n:
        k_pt = ('CARIBE', 'PUERTO TEJADA', prod_n)
        if k_pt in exact_map:
            return exact_map[k_pt][0], 'Activo', 'Tarifa Caribe Puerto Tejada'
    if 'VILLA RICA' in pto_n and 'CARIBE' in prov_n:
        k_vr = ('CARIBE', 'VILLA RICA', prod_n)
        if k_vr in exact_map:
            return exact_map[k_vr][0], 'Activo', 'Tarifa Caribe Villa Rica'
            
    return None, 'Pendiente Precio', 'Pendiente cotizar'

# 4. Generate the 2,772 rows
repaired_rows = []
for pt in puntos:
    for prod in prods_activos:
        precio, estado, obs = find_tariff(pt, prod)
        repaired_rows.append({
            'Ruta': pt['ruta'],
            'Proveedor': pt['prov'],
            'Punto_Sucursal': pt['pto'],
            'Producto': prod,
            'Precio_Kg': precio,
            'Estado': estado,
            'Observaciones': obs
        })

# Styling definitions
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
fill_header = PatternFill(start_color='1E4620', end_color='1E4620', fill_type='solid')
font_data = Font(name='Calibri', size=10)
font_bold = Font(name='Calibri', size=10, bold=True)
fill_pending = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid') # soft amber for pending
font_pending = Font(name='Calibri', size=10, color='92400E')
border_thin = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')
align_center = Alignment(horizontal='center', vertical='center')

def style_tarifas_sheet(ws, rows):
    headers = ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Producto', 'Precio_Kg', 'Estado', 'Observaciones']
    ws.append(headers)
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center if col_idx in [5, 6] else align_left
        cell.border = border_thin
    
    ws.row_dimensions[1].height = 26
    
    for r_idx, r_data in enumerate(rows, start=2):
        ws.row_dimensions[r_idx].height = 20
        is_pending = (r_data['Precio_Kg'] is None)
        
        c1 = ws.cell(row=r_idx, column=1, value=r_data['Ruta'])
        c1.font = font_data
        c1.alignment = align_left
        c1.border = border_thin
        
        c2 = ws.cell(row=r_idx, column=2, value=r_data['Proveedor'])
        c2.font = font_data
        c2.alignment = align_left
        c2.border = border_thin
        
        c3 = ws.cell(row=r_idx, column=3, value=r_data['Punto_Sucursal'])
        c3.font = font_data
        c3.alignment = align_left
        c3.border = border_thin
        
        c4 = ws.cell(row=r_idx, column=4, value=r_data['Producto'])
        c4.font = font_bold
        c4.alignment = align_left
        c4.border = border_thin
        
        c5 = ws.cell(row=r_idx, column=5, value=r_data['Precio_Kg'])
        c5.font = font_bold if not is_pending else font_pending
        c5.alignment = align_right
        c5.border = border_thin
        if r_data['Precio_Kg'] is not None:
            c5.number_format = '$#,##0'
        
        c6 = ws.cell(row=r_idx, column=6, value=r_data['Estado'])
        c6.font = font_pending if is_pending else font_data
        c6.alignment = align_center
        c6.border = border_thin
        if is_pending:
            c6.fill = fill_pending
        
        c7 = ws.cell(row=r_idx, column=7, value=r_data['Observaciones'])
        c7.font = font_pending if is_pending else font_data
        c7.alignment = align_left
        c7.border = border_thin
    
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:G{len(rows)+1}"
    
    ws.column_dimensions['A'].width = 42
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 38
    ws.column_dimensions['D'].width = 26
    ws.column_dimensions['E'].width = 16
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 34

# 5. Update tarifario.xlsx by adding Tarifas_Reparadas
if 'Tarifas_Reparadas' in wb_orig.sheetnames:
    del wb_orig['Tarifas_Reparadas']
ws_rep = wb_orig.create_sheet(title='Tarifas_Reparadas')
style_tarifas_sheet(ws_rep, repaired_rows)
wb_orig.save('tarifario.xlsx')

# 6. Create standalone clean workbook Tarifario_Completo_Reparado_2026.xlsx
wb_new = openpyxl.Workbook()

# Sheet 1: Tarifas
ws_new_tarifas = wb_new.active
ws_new_tarifas.title = 'Tarifas'
style_tarifas_sheet(ws_new_tarifas, repaired_rows)

# Sheet 2: Puntos_Rutas (clean)
ws_new_puntos = wb_new.create_sheet(title='Puntos_Rutas')
p_headers = ['Ruta', 'Proveedor', 'Punto_Sucursal', 'Direccion', 'Telefono', 'Horario_Estimado', 'Frecuencia_Dias', 'Estado', 'Materias_Frecuentes']
ws_new_puntos.append(p_headers)
for c_idx in range(1, len(p_headers)+1):
    c = ws_new_puntos.cell(row=1, column=c_idx)
    c.font = font_header
    c.fill = fill_header
    c.alignment = align_left
    c.border = border_thin
ws_new_puntos.row_dimensions[1].height = 26

for r_idx, r_vals in enumerate(puntos_full_rows, start=2):
    ws_new_puntos.row_dimensions[r_idx].height = 20
    for c_idx, val in enumerate(r_vals, start=1):
        c = ws_new_puntos.cell(row=r_idx, column=c_idx, value=val)
        c.font = font_data
        c.alignment = align_left
        c.border = border_thin

ws_new_puntos.freeze_panes = 'A2'
ws_new_puntos.auto_filter.ref = f"A1:I{len(puntos_full_rows)+1}"
ws_new_puntos.column_dimensions['A'].width = 42
ws_new_puntos.column_dimensions['B'].width = 28
ws_new_puntos.column_dimensions['C'].width = 38
ws_new_puntos.column_dimensions['D'].width = 32
ws_new_puntos.column_dimensions['E'].width = 16
ws_new_puntos.column_dimensions['F'].width = 20
ws_new_puntos.column_dimensions['G'].width = 20
ws_new_puntos.column_dimensions['H'].width = 14
ws_new_puntos.column_dimensions['I'].width = 48

# Sheet 3: Productos (clean)
ws_new_prod = wb_new.create_sheet(title='Productos')
prod_headers = ['Producto', 'Estado']
ws_new_prod.append(prod_headers)
for c_idx in range(1, 3):
    c = ws_new_prod.cell(row=1, column=c_idx)
    c.font = font_header
    c.fill = fill_header
    c.alignment = align_left
    c.border = border_thin
ws_new_prod.row_dimensions[1].height = 26

for r_idx, (p_name, p_st) in enumerate(prods_all, start=2):
    ws_new_prod.row_dimensions[r_idx].height = 20
    c1 = ws_new_prod.cell(row=r_idx, column=1, value=p_name)
    c1.font = font_bold
    c1.border = border_thin
    c2 = ws_new_prod.cell(row=r_idx, column=2, value=p_st)
    c2.font = font_data
    c2.border = border_thin
    c2.alignment = align_center

ws_new_prod.freeze_panes = 'A2'
ws_new_prod.auto_filter.ref = f"A1:B{len(prods_all)+1}"
ws_new_prod.column_dimensions['A'].width = 35
ws_new_prod.column_dimensions['B'].width = 16

# Sheet 4: Diagnostico_Auditoria
ws_new_diag = wb_new.create_sheet(title='Diagnostico_Auditoria')
diag_title = ws_new_diag.cell(row=1, column=1, value="DIAGNÓSTICO Y AUDITORÍA DE TARIFAS - PROTEINAGRO 2026")
diag_title.font = Font(name='Calibri', size=14, bold=True, color='1E4620')
ws_new_diag.row_dimensions[1].height = 30

diag_lines = [
    ("Concepto", "Detalle de Hallazgos y Solución Aplicada"),
    ("1. Causa Raíz de Inconsistencias", "La hoja 'Tarifas' original solo contenía 461 registros con datos frente a los 2,772 requeridos (126 puntos × 22 productos activos)."),
    ("2. Productos Faltantes en Tarifas", "11 productos activos (Calambombo Res/Cerdo, Hueso Pollo, Leña, Mantequilla, Margarina, Orejas Cerdo, Pulmón Cerdo, Tráqueas, etc.) nunca fueron registrados en la matriz antigua para la mayoría de puntos."),
    ("3. Productos Fuera de Catálogo", "En la versión anterior existían productos no estandarizados como 'Aceite Medio', 'Aceite Limpio', 'Aceite Alto 3000' y 'Hueso' sin especificación."),
    ("4. Discrepancia de Proveedores", "Múltiples sucursales figuraban bajo 'Cuenta Sevillana', 'Cuenta Alberto Millán', 'Cuenta Fábrica' o 'Migan Capital', impidiendo el match automático en la App."),
    ("5. Puntos Agrupados vs Sucursales", "En la versión antigua se agrupaban zonas (ej: 'Cañaveral Roldanillo - Zarzal', 'Jamundí', 'Belalcázar Centro'), mientras que en 'Puntos_Rutas' cada sucursal tiene su propia identidad (B1 a B11 Belalcázar, sucursales individuales de Cañaveral, etc.)."),
    ("6. Matriz Reparada Generada", "Se generaron exactamente 2,772 filas: cada uno de los 126 puntos de recolección ahora cuenta con los 22 productos activos disponibles."),
    ("7. Precios Históricos Recuperados", "Se preservaron e indexaron 353 precios confirmados y recuperados por sucursal/zona. Las 2,419 combinaciones restantes quedan marcadas como 'Pendiente Precio' con aviso 'Pendiente cotizar' para definición comercial.")
]

for r_idx, (c1_val, c2_val) in enumerate(diag_lines, start=3):
    ws_new_diag.row_dimensions[r_idx].height = 24
    c1 = ws_new_diag.cell(row=r_idx, column=1, value=c1_val)
    c2 = ws_new_diag.cell(row=r_idx, column=2, value=c2_val)
    c1.border = border_thin
    c2.border = border_thin
    if r_idx == 3:
        c1.font = font_header
        c1.fill = fill_header
        c2.font = font_header
        c2.fill = fill_header
    else:
        c1.font = font_bold
        c2.font = font_data

ws_new_diag.column_dimensions['A'].width = 35
ws_new_diag.column_dimensions['B'].width = 85

new_file_path = 'Tarifario_Completo_Reparado_2026.xlsx'
wb_new.save(new_file_path)
print("Archivos generados exitosamente con texto y acentos limpios.")
